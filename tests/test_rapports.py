"""UC23 · Générer un rapport de performance — UC24 · Exporter un rapport. Tests
d'intégration (base PostgreSQL de test)."""

from __future__ import annotations

from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from app.bd.connexion import transaction
from app.bd.depots.alertes import DepotAlertes
from app.bd.depots.historique import DepotHistorique
from app.bd.depots.modeles import DepotModeles
from app.bd.depots.plans import DepotPlansCharge
from app.bd.depots.previsions_ressources import DepotPrevisionsRessources
from app.config import DOSSIER_RAPPORTS
from app.erreurs import AccesRefuse, OperationImpossible
from app.services import admin, auth, modeles, rapports
from app.utils.dates import jours_semaine, lundi_de
from tests.test_modeles import _inserer_historique

pytestmark = pytest.mark.integration

DATE_RAPPORT = date(2026, 2, 2)  # dans les 100 jours d'historique déterministe de test


@pytest.fixture
def site_zone(ctx_admin):
    site_id = admin.enregistrer_site(ctx_admin, "Plateforme Test")
    zone_id = admin.enregistrer_zone(ctx_admin, site_id, "Réception", "chariot_elevateur", 7.5)
    return site_id, zone_id


@pytest.fixture
def ctx_responsable(site_zone):
    from app.bd.depots.utilisateurs import DepotUtilisateurs

    site_id, _zone_id = site_zone
    with transaction() as cur:
        h, s = auth.hacher_mot_de_passe("Resp2026!")
        depot = DepotUtilisateurs(cur)
        uid = depot.creer("resp", "R", "R", "", h, s, "responsable")
        depot.definir_sites(uid, [site_id])
    return auth.authentifier("resp", "Resp2026!")


@pytest.fixture
def ctx_planificateur(site_zone):
    from app.bd.depots.utilisateurs import DepotUtilisateurs

    site_id, _zone_id = site_zone
    with transaction() as cur:
        h, s = auth.hacher_mot_de_passe("Planif2026!")
        depot = DepotUtilisateurs(cur)
        uid = depot.creer("planif", "P", "P", "", h, s, "planificateur")
        depot.definir_sites(uid, [site_id])
    return auth.authentifier("planif", "Planif2026!")


def _semer_historique_et_alerte(site_id: int, zone_id: int, date_reference: date) -> None:
    """Heures sup à 20 % du travaillé (TAUX_HS rouge, cible générale 5 %) sur la semaine, et
    une alerte concernant le même lundi."""
    lundi = lundi_de(date_reference)
    with transaction() as cur:
        DepotHistorique(cur).upsert_plusieurs(
            [
                {
                    "site_id": site_id,
                    "zone_id": zone_id,
                    "date_jour": jour,
                    "volume_traite": 1000.0,
                    "effectif_present": 10,
                    "heures_travaillees": 100.0,
                    "heures_sup": 20.0,
                    "heures_interim": 0,
                    "heures_absence": 0,
                    "heures_inactives": 0,
                    "equipements_mobilises": 2,
                    "heures_usage_equipement": 20.0,
                    "heures_disponibles_equipement": 24.0,
                    "heures_panne_equipement": 0.0,
                    "cout_rh": 100.0,
                    "commandes_a_temps": 95,
                    "commandes_totales": 100,
                    "indicateur_pic": False,
                    "source": "test",
                }
                for jour in jours_semaine(lundi)
            ]
        )
        DepotAlertes(cur).emettre(
            "penurie_equipement",
            "orange",
            None,
            site_id,
            zone_id,
            lundi,
            f"test_penurie_{site_id}_{zone_id}",
            "Pénurie d'équipements de test.",
        )


def test_generer_rapport_refuse_hors_responsable(ctx_planificateur, site_zone):
    site_id, _zone_id = site_zone
    with pytest.raises(AccesRefuse):
        rapports.generer_rapport(ctx_planificateur, site_id, "semaine")


def test_generer_rapport_refuse_format_invalide(ctx_responsable, site_zone):
    site_id, _zone_id = site_zone
    with pytest.raises(ValueError):
        rapports.generer_rapport(ctx_responsable, site_id, "semaine", format_rapport="doc")


def test_generer_rapport_pdf_excel_ecrit_les_deux_fichiers(ctx_responsable, site_zone):
    site_id, zone_id = site_zone
    date_reference = date(2026, 3, 18)  # mercredi
    _semer_historique_et_alerte(site_id, zone_id, date_reference)

    resultat = rapports.generer_rapport(
        ctx_responsable, site_id, "semaine", date_reference, "pdf_excel"
    )
    assert resultat["chemin_pdf"] is not None
    assert resultat["chemin_excel"] is not None
    assert (DOSSIER_RAPPORTS / resultat["chemin_pdf"]).is_file()
    assert (DOSSIER_RAPPORTS / resultat["chemin_excel"]).is_file()

    taux_hs = next(v for v in resultat["kpi"] if v["kpi_code"] == "TAUX_HS")
    assert taux_hs["valeur"] == pytest.approx(20.0)
    assert taux_hs["statut"] == "rouge"
    assert len(resultat["alertes"]) == 1
    assert resultat["alertes"][0]["type"] == "penurie_equipement"

    with transaction() as cur:
        cur.execute("SELECT contenu FROM rapports WHERE id = %s", (resultat["rapport_id"],))
        contenu = cur.fetchone()["contenu"]
    assert contenu["site"] == "Plateforme Test"
    assert len(contenu["kpi"]) == len(resultat["kpi"])


def _preparer_rapprochement_et_plan(ctx_admin, site_id: int, zone_id: int) -> None:
    """Historique entraînable, une prévision RL rapprochée au réel, et un plan de charge
    validé, tous sur :data:`DATE_RAPPORT` (sections « Réel/RL/RN », « Comparaison RL vs RN »
    et « Plan de charge validé et écarts » du rapport, UC23). Entraîner et activer un modèle
    (UC08, UC10) sont réservés à l'administrateur ; le reste (dépôts directs) est indifférent
    au rôle."""
    _inserer_historique(site_id, zone_id)
    resume = modeles.entrainer_modeles(ctx_admin, site_id, zone_id)
    for version in resume.versions:
        modeles.activer_version(
            ctx_admin,
            version.version_id,
            retenir_pour_plan=version.methode == "regression_lineaire",
        )
    with transaction() as cur:
        depot_modeles = DepotModeles(cur)
        version = depot_modeles.version_active(site_id, zone_id, "regression_lineaire", "heures")
        version_eqp = depot_modeles.version_active(
            site_id, zone_id, "regression_lineaire", "equipements"
        )
        DepotPrevisionsRessources(cur).inserer_plusieurs(
            [
                {
                    "site_id": site_id,
                    "zone_id": zone_id,
                    "date_jour": DATE_RAPPORT,
                    "modele_version_id": version["id"],
                    "modele_version_equipements_id": version_eqp["id"],
                    "methode": "regression_lineaire",
                    "volume_prevu": 1000.0,
                    "heures": 60.0,
                    "effectif": 8,
                    "equipements": 3,
                    "ic_bas": 55.0,
                    "ic_haut": 65.0,
                    "ic_bas_equipements": 1,
                    "ic_haut_equipements": 5,
                    "date_generation": datetime.now(),
                }
            ]
        )
        depot_plans = DepotPlansCharge(cur)
        plan_id = depot_plans.creer_plan(
            site_id, lundi_de(DATE_RAPPORT), "regression_lineaire", None
        )
        depot_plans.inserer_lignes(
            plan_id,
            [
                {
                    "zone_id": zone_id,
                    "date_jour": DATE_RAPPORT,
                    "besoin_heures": 60.0,
                    "besoin_effectif": 10,
                    "besoin_equipements": 3,
                    "effectif_planifie": 6,
                    "interim_planifie": 0,
                    "equipements_planifies": 3,
                    "capacite_effectif": 12,
                    "capacite_equipements": 4,
                    "commentaire": "",
                }
            ],
        )
        depot_plans.mettre_a_jour_plan(plan_id, statut="valide")


def test_generer_rapport_couvre_les_sept_sections_uc23(ctx_admin, ctx_responsable, site_zone):
    site_id, zone_id = site_zone
    _preparer_rapprochement_et_plan(ctx_admin, site_id, zone_id)

    resultat = rapports.generer_rapport(ctx_responsable, site_id, "jour", DATE_RAPPORT, "pdf_excel")

    with transaction() as cur:
        cur.execute("SELECT contenu FROM rapports WHERE id = %s", (resultat["rapport_id"],))
        contenu = cur.fetchone()["contenu"]
    # 2. synthèse, en 5 lignes maximum.
    assert 1 <= len(contenu["synthese"]) <= 5

    classeur = load_workbook(DOSSIER_RAPPORTS / resultat["chemin_excel"])
    # Une feuille par section tabulaire (3. KPI, 5. comparaison RL/RN, 6. alertes,
    # 7. plan de charge), plus la synthèse (2.) et l'équivalent tabulaire du graphique (4.).
    assert set(classeur.sheetnames) == {
        "Synthèse",
        "KPI",
        "Réel RL RN",
        "Comparaison RL RN",
        "Alertes",
        "Plan de charge",
    }
    feuille_comparaison = classeur["Comparaison RL RN"]
    lignes_comparaison = list(feuille_comparaison.iter_rows(values_only=True))
    assert lignes_comparaison[0] == ("Métrique", "RL", "RN", "Meilleure méthode")
    assert len(lignes_comparaison) == 7  # en-tête + 6 métriques (5 formules + taux de victoire)

    feuille_plan = classeur["Plan de charge"]
    lignes_plan = list(feuille_plan.iter_rows(values_only=True))
    assert len(lignes_plan) == 2  # en-tête + la ligne du plan validé
    assert lignes_plan[1][0] == "Réception"  # zone
    assert lignes_plan[1][4] == "-4"  # écart effectif : 6 planifiés - 10 besoin

    feuille_reel = classeur["Réel RL RN"]
    lignes_reel = list(feuille_reel.iter_rows(values_only=True))
    assert len(lignes_reel) == 2  # en-tête + un jour de rapprochement


def test_generer_rapport_pdf_seul_n_ecrit_pas_excel(ctx_responsable, site_zone):
    site_id, _zone_id = site_zone
    resultat = rapports.generer_rapport(
        ctx_responsable, site_id, "semaine", date(2026, 3, 18), "pdf"
    )
    assert resultat["chemin_pdf"] is not None
    assert resultat["chemin_excel"] is None


def test_lister_rapports_plus_recent_d_abord(ctx_responsable, site_zone):
    site_id, _zone_id = site_zone
    rapports.generer_rapport(ctx_responsable, site_id, "semaine", date(2026, 3, 18), "excel")
    rapports.generer_rapport(ctx_responsable, site_id, "semaine", date(2026, 3, 25), "excel")
    liste = rapports.lister_rapports(ctx_responsable, site_id)
    assert len(liste) == 2
    assert liste[0]["date_debut"] > liste[1]["date_debut"]
    assert liste[0]["genere_par_identifiant"] == "resp"
    assert liste[0]["genere_par_systeme"] is False


def test_exporter_rapport_pdf_et_excel(ctx_responsable, site_zone):
    site_id, _zone_id = site_zone
    resultat = rapports.generer_rapport(
        ctx_responsable, site_id, "semaine", date(2026, 3, 18), "pdf_excel"
    )
    chemin_pdf = rapports.exporter_rapport(ctx_responsable, resultat["rapport_id"], "pdf")
    assert chemin_pdf.is_file()
    assert chemin_pdf.suffix == ".pdf"
    chemin_excel = rapports.exporter_rapport(ctx_responsable, resultat["rapport_id"], "excel")
    assert chemin_excel.is_file()
    assert chemin_excel.suffix == ".xlsx"


def test_exporter_rapport_format_absent_refuse(ctx_responsable, site_zone):
    site_id, _zone_id = site_zone
    resultat = rapports.generer_rapport(
        ctx_responsable, site_id, "semaine", date(2026, 3, 18), "pdf"
    )
    with pytest.raises(OperationImpossible):
        rapports.exporter_rapport(ctx_responsable, resultat["rapport_id"], "excel")


def test_exporter_rapport_introuvable(ctx_responsable):
    with pytest.raises(OperationImpossible):
        rapports.exporter_rapport(ctx_responsable, 999999, "pdf")


def test_exporter_rapport_refuse_hors_responsable(ctx_planificateur, ctx_responsable, site_zone):
    site_id, _zone_id = site_zone
    resultat = rapports.generer_rapport(
        ctx_responsable, site_id, "semaine", date(2026, 3, 18), "pdf"
    )
    with pytest.raises(AccesRefuse):
        rapports.exporter_rapport(ctx_planificateur, resultat["rapport_id"], "pdf")
