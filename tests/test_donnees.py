"""UC04 · Importer / saisir l'historique — UC05 · même chose pour les prévisions de volume."""

from datetime import date, timedelta

import pytest

from app.bd.connexion import transaction
from app.erreurs import DonneesInvalides
from app.services import donnees

pytestmark = pytest.mark.integration


@pytest.fixture
def site_zone(ctx_admin):
    from app.services import admin

    site_id = admin.enregistrer_site(ctx_admin, "Plateforme Test")
    zone_id = admin.enregistrer_zone(ctx_admin, site_id, "Réception", "chariot_elevateur", 7.5)
    return site_id, zone_id


@pytest.fixture
def ctx_planificateur(bd_vierge, site_zone):
    from app.bd.depots.utilisateurs import DepotUtilisateurs
    from app.services import auth

    site_id, _ = site_zone
    with transaction() as cur:
        hash_mdp, sel = auth.hacher_mot_de_passe("Planif2026!")
        depot = DepotUtilisateurs(cur)
        uid = depot.creer("planif", "Test", "Plan", "", hash_mdp, sel, "planificateur")
        depot.definir_sites(uid, [site_id])
    return auth.authentifier("planif", "Planif2026!")


def _recent(jours_avant: int = 3) -> str:
    """Date récente (dans la fenêtre des 60 derniers jours), au format JJ/MM/AAAA."""
    return (date.today() - timedelta(days=jours_avant)).strftime("%d/%m/%Y")


def ligne_historique(**remplacements) -> dict:
    base = {
        "date": _recent(),
        "volume_traite": "1250,5",
        "effectif_present": "10",
        "heures_travaillees": "75",
        "heures_sup": "5",
        "heures_interim": "0",
        "heures_absence": "0",
        "heures_inactives": "3",
        "equipements_mobilises": "4",
        "heures_usage_equipement": "28",
        "heures_disponibles_equipement": "30",
        "heures_panne_equipement": "0",
        "cout_rh": "3487,50",
        "commandes_a_temps": "118",
        "commandes_totales": "120",
        "indicateur_pic": "Non",
    }
    base.update(remplacements)
    return base


# ---------------------------------------------------------------------
# UC04 · Saisie manuelle
# ---------------------------------------------------------------------
def test_saisie_historique_creation_et_mise_a_jour(ctx_planificateur, site_zone):
    site_id, zone_id = site_zone
    avertissements = donnees.saisir_historique(
        ctx_planificateur, site_id, zone_id, ligne_historique()
    )
    assert avertissements == []
    lignes = donnees.lister_historique(ctx_planificateur, site_id)
    assert len(lignes) == 1 and lignes[0]["volume_traite"] == 1250.5

    # Une deuxième saisie à la même date met à jour (ON CONFLICT DO UPDATE), pas de doublon.
    donnees.saisir_historique(
        ctx_planificateur, site_id, zone_id, ligne_historique(volume_traite="1300")
    )
    lignes = donnees.lister_historique(ctx_planificateur, site_id)
    assert len(lignes) == 1 and lignes[0]["volume_traite"] == 1300.0
    assert lignes[0]["source"] == "saisie"


def test_saisie_historique_invalide(ctx_planificateur, site_zone):
    site_id, zone_id = site_zone
    with pytest.raises(DonneesInvalides) as erreur:
        donnees.saisir_historique(
            ctx_planificateur, site_id, zone_id, ligne_historique(volume_traite=None)
        )
    assert erreur.value.erreurs == {"volume_traite": "Le champ « Volume traité » est obligatoire."}
    assert donnees.lister_historique(ctx_planificateur, site_id) == []


def test_saisie_historique_refuse_hors_site(ctx_planificateur, site_zone):
    from app.erreurs import AccesRefuse

    with pytest.raises(AccesRefuse):
        donnees.saisir_historique(ctx_planificateur, 9999, 1, ligne_historique())


def test_saisie_historique_avertissement_remonte(ctx_planificateur, site_zone):
    site_id, zone_id = site_zone
    lundi = date(2026, 3, 2)
    for i in range(8):
        donnees.saisir_historique(
            ctx_planificateur,
            site_id,
            zone_id,
            ligne_historique(
                date=(lundi + timedelta(weeks=i)).strftime("%d/%m/%Y"), volume_traite="1000"
            ),
        )
    avertissements = donnees.saisir_historique(
        ctx_planificateur,
        site_id,
        zone_id,
        ligne_historique(
            date=(lundi + timedelta(weeks=8)).strftime("%d/%m/%Y"), volume_traite="9000"
        ),
    )
    assert len(avertissements) == 1 and "inhabituelle" in avertissements[0]


# ---------------------------------------------------------------------
# UC04 · Import de fichier
# ---------------------------------------------------------------------
def _ecrire_csv(tmp_path, nom, entetes, lignes):
    chemin = tmp_path / nom
    contenu = ";".join(entetes) + "\n" + "\n".join(";".join(l) for l in lignes)
    chemin.write_text(contenu, encoding="utf-8")
    return chemin


def test_import_historique_valide_et_rejete(ctx_planificateur, site_zone, tmp_path):
    site_id, zone_id = site_zone
    entetes = ["site", "zone", "date", "volume_traite", "effectif_present", "heures_travaillees"]
    chemin = _ecrire_csv(
        tmp_path,
        "historique.csv",
        entetes,
        [
            ["Plateforme Test", "Réception", _recent(1), "1250,5", "10", "75"],
            ["Plateforme Test", "Zone inconnue", _recent(2), "1000", "8", "60"],
            ["Plateforme Test", "Réception", _recent(3), "", "8", "60"],
        ],
    )
    resultat = donnees.importer_historique(ctx_planificateur, chemin)
    assert len(resultat.valides) == 1
    assert len(resultat.rejets) == 2
    assert donnees.lister_historique(ctx_planificateur, site_id) == []  # pas encore écrit

    nombre = donnees.enregistrer_lignes_historique(ctx_planificateur, resultat.valides)
    assert nombre == 1
    lignes = donnees.lister_historique(ctx_planificateur, site_id)
    assert len(lignes) == 1 and lignes[0]["source"] == "import"


def test_import_historique_doublon_dans_fichier(ctx_planificateur, site_zone, tmp_path):
    entetes = ["site", "zone", "date", "volume_traite", "effectif_present", "heures_travaillees"]
    chemin = _ecrire_csv(
        tmp_path,
        "historique.csv",
        entetes,
        [
            ["Plateforme Test", "Réception", "09/03/2026", "1250,5", "10", "75"],
            ["Plateforme Test", "Réception", "09/03/2026", "1300", "10", "75"],
        ],
    )
    resultat = donnees.importer_historique(ctx_planificateur, chemin)
    assert len(resultat.valides) == 1 and len(resultat.rejets) == 1
    assert resultat.rejets[0].ligne == 2


def test_import_fichier_inexistant(ctx_planificateur, tmp_path):
    with pytest.raises(DonneesInvalides, match="introuvable"):
        donnees.importer_historique(ctx_planificateur, tmp_path / "absent.csv")


def test_import_format_non_pris_en_charge(ctx_planificateur, tmp_path):
    chemin = tmp_path / "donnees.txt"
    chemin.write_text("site;zone\n")
    with pytest.raises(DonneesInvalides, match="Format de fichier non pris en charge"):
        donnees.importer_historique(ctx_planificateur, chemin)


# ---------------------------------------------------------------------
# UC05 · Prévisions de volume
# ---------------------------------------------------------------------
def test_saisie_prevision_volume(ctx_planificateur, site_zone):
    site_id, zone_id = site_zone
    demain = (date.today() + timedelta(days=1)).strftime("%d/%m/%Y")
    donnees.saisir_prevision_volume(
        ctx_planificateur,
        site_id,
        zone_id,
        {"date": demain, "volume_prevu": "1750", "indicateur_pic": "Oui"},
    )
    lignes = donnees.lister_previsions_volume(ctx_planificateur, site_id)
    assert len(lignes) == 1 and lignes[0]["volume_prevu"] == 1750.0
    assert lignes[0]["indicateur_pic"] is True


def test_saisie_prevision_date_passee_refusee(ctx_planificateur, site_zone):
    site_id, zone_id = site_zone
    hier = (date.today() - timedelta(days=1)).strftime("%d/%m/%Y")
    with pytest.raises(DonneesInvalides) as erreur:
        donnees.saisir_prevision_volume(
            ctx_planificateur, site_id, zone_id, {"date": hier, "volume_prevu": "1000"}
        )
    assert "passée" in erreur.value.erreurs["date"]


def test_import_previsions_volume(ctx_planificateur, site_zone, tmp_path):
    site_id, zone_id = site_zone
    demain = (date.today() + timedelta(days=1)).strftime("%d/%m/%Y")
    chemin = _ecrire_csv(
        tmp_path,
        "previsions.csv",
        ["site", "zone", "date", "volume_prevu"],
        [["Plateforme Test", "Réception", demain, "1750"]],
    )
    resultat = donnees.importer_previsions_volume(ctx_planificateur, chemin)
    assert len(resultat.valides) == 1 and not resultat.rejets
    assert donnees.enregistrer_lignes_previsions(ctx_planificateur, resultat.valides) == 1
    assert donnees.lister_previsions_volume(ctx_planificateur, site_id)[0]["source"] == "import"


# ---------------------------------------------------------------------
# Modèles de fichiers et rapport d'erreurs
# ---------------------------------------------------------------------
def test_modeles_de_fichiers_existent():
    for cle in ("historique", "prevision"):
        chemin = donnees.chemin_modele_fichier(cle)
        assert chemin.is_file()


def test_export_rapport_erreurs(tmp_path):
    from app.services.qualite import Avertissement, Rejet

    chemin = tmp_path / "rapport.xlsx"
    donnees.exporter_rapport_erreurs(
        chemin,
        [Rejet(2, "volume_traite", "Le champ « Volume traité » est obligatoire.")],
        [Avertissement(3, "volume_traite", "Valeur inhabituelle.")],
    )
    assert chemin.is_file()
    from openpyxl import load_workbook

    classeur = load_workbook(chemin)
    assert classeur.sheetnames == ["Rejets", "Avertissements"]
    assert classeur["Rejets"]["A2"].value == 2
